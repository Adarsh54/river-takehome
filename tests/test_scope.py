import copy
import json
import hashlib
import threading
from pathlib import Path
import pytest
import requests
from http.server import ThreadingHTTPServer
from apex_benchmark.metrics import validate_metrics,metric_labels,extract_custom,augment_custom,definition_hash
from apex_benchmark.fund_browser import FundBrowser,COMPANIES,FUNDS,NAMES
from apex_benchmark.judgment import build_packet,validate_decision,JudgmentError
from apex_benchmark.workspace import Workspace,handler
from apex_benchmark.word_report import export_word, export_evidence_word, committee_summary
from apex_benchmark.resolve import resolve_metric

ROOT=Path(__file__).resolve().parents[1]
CONFIG=json.loads((ROOT/'config/engagement.json').read_text())
M={'id':'custom_turnover','label':'Portfolio turnover','definition':'Reported annual portfolio turnover percentage, excluding originations.','unit':'%','scope':'fund','period_basis':'fiscal_year','keywords':['portfolio turnover'],'calculation':'reported_only'}
F={'ticker':'TEST','name':'Test Fund','cik':'0000001234','structure':'bdc','share_class':'Common','return_basis':'shareholder'}
FILING={'ticker':'TEST','form':'10-K','filingDate':'2025-02-01','reportDate':'2024-12-31','accessionNumber':'acc'}
BODY=b'<table><tr><th>Year ended December 31 2024</th></tr><tr><td>Portfolio turnover</td><td>24.5%</td></tr></table>'
SOURCE={'url':'https://www.sec.gov/Archives/example.htm','sha256':hashlib.sha256(BODY).hexdigest()}


def test_metric_definition_validation_and_no_regex_execution():
 c=copy.deepcopy(CONFIG);c['custom_metrics']=[M];validate_metrics(c)
 assert 'custom_turnover' in metric_labels(c)
 for bad in [dict(M,id='nav'),dict(M,keywords=[]),dict(M,calculation='eval(1)'),dict(M,scope='anything')]:
  with pytest.raises(ValueError):validate_metrics(dict(c,custom_metrics=[bad]))
 assert not extract_custom(BODY,F,FILING,SOURCE,[dict(M,keywords=['.*'])])


def test_custom_extraction_keeps_table_headers_and_definition_identity():
 cs=extract_custom(BODY,F,FILING,SOURCE,[M]);assert len(cs)==1
 c=cs[0];assert 'Year ended December 31 2024' in c['quote']
 assert c['value'] is None and not c['eligible']
 assert c['unit']=='%' and c['extraction_meta']['metric_definition_sha256']==definition_hash(M)


def custom_packet():
 cfg=dict(CONFIG,funds=[F],custom_metrics=[M]);cs=extract_custom(BODY,F,FILING,SOURCE,[M])
 result=resolve_metric(cs,F,M['id'],cfg)
 return build_packet({'config':cfg,'filings':[FILING],'candidates':cs},result)


def test_custom_llm_requires_definition_scope_and_reported_value():
 p=custom_packet();c=p['evidence'][0]
 d={'decision':'publish','selected_candidate_id':c['id'],'value':24.5,'period':'2024-12-31','unit':'%',
 'basis':'custom:custom_turnover','share_class':'Fund','directly_reported':True,'rationale':'Annual turnover is reported explicitly.','comparability':'Issuer definition retained.','unresolved_issues':[],
 'citations':[{'candidate_id':c['id'],'quote':c['quote']}],'assessments':[{'candidate_id':c['id'],'reason':'Matching reported fiscal-year turnover.'}]}
 assert validate_decision(d,p)
 for change in [dict(unit='bps'),dict(basis='unconfirmed'),dict(share_class='Common'),dict(value=245),dict(directly_reported=False)]:
  with pytest.raises(JudgmentError):validate_decision(dict(d,**change),p)
 changed=copy.deepcopy(p);changed['custom_metric_definition']['definition']='Different denominator'
 assert json.dumps(changed,sort_keys=True)!=json.dumps(p,sort_keys=True)


class SecFixture:
 def json(self,url):
  if url==COMPANIES:return {'0':{'cik_str':1234,'title':'Test Capital','ticker':'TEST'}}
  if url==FUNDS:return {'fields':['cik','seriesId','classId','symbol'],'data':[[5678,'S123','C123','CLASSX']]}
  return {'name':'Actual SEC name','tickers':['TEST'],'filings':{'recent':{'form':['10-K','10-Q','4']}}}
 def get(self,url):
  assert url==NAMES
  return b'Test Capital:1234:\nUnlisted Fund:5678:\n',{}


def test_sec_browser_name_ticker_cik_and_share_class():
 b=FundBrowser(SecFixture())
 assert b.search('TEST')['results'][0]['cik']=='0000001234'
 r=b.search('CLASSX')['results'][0];assert r['cik']=='0000005678' and r['classes'][0]['seriesId']=='S123'
 assert b.search('Unlisted')['results'][0]['name']=='Unlisted Fund'
 assert b.search('1234')['results'][0]['supported_forms']==['10-K','10-Q']
 with pytest.raises(ValueError):b.detail('../etc')


def test_workspace_persistence_and_csrf(tmp_path,monkeypatch):
 (tmp_path/'config').mkdir();(tmp_path/'config/engagement.json').write_text(json.dumps(CONFIG))
 w=Workspace(tmp_path,tmp_path/'config/workspace.json',tmp_path/'report');w.browser=FundBrowser(SecFixture())
 added=w.add_fund(dict(F,ticker='TEST'));assert added['name']=='Actual SEC name' and len(w.config['funds'])==5
 assert json.loads(w.config_path.read_text())['funds'][-1]['cik']=='0000001234'
 server=ThreadingHTTPServer(('127.0.0.1',0),handler(w));threading.Thread(target=server.serve_forever,daemon=True).start()
 url=f'http://127.0.0.1:{server.server_port}'
 try:
  assert requests.post(url+'/api/metrics',json={'metrics':[M]}).status_code==403
  assert requests.get(url+'/api/state',headers={'Host':'evil.example'}).status_code==403
  token=requests.get(url+'/api/state').json()['csrf']
  assert requests.post(url+'/api/metrics',json={'metrics':[M]},headers={'X-CSRF-Token':token}).status_code==200
  assert w.config['custom_metrics']==[M]
  assert requests.get(url+'/report/.env').status_code==404
 finally:server.shutdown();server.server_close()


def test_custom_source_integrity_and_change_reextraction(tmp_path):
 (tmp_path/SOURCE['sha256']).write_bytes(BODY)
 run={'config':dict(CONFIG,funds=[F],custom_metrics=[M]),'filings':[FILING],'documents':[dict(SOURCE,ticker='TEST',accession='acc')],'candidates':[],'errors':[]}
 augment_custom(run,tmp_path);assert len(run['candidates'])==1
 run['config']['custom_metrics']=[dict(M,keywords=['no match'])];augment_custom(run,tmp_path);assert not run['candidates']
 (tmp_path/SOURCE['sha256']).write_bytes(b'changed');augment_custom(run,tmp_path);assert run['errors']


def test_word_document_contains_dynamic_funds_custom_metrics_and_source_links(tmp_path):
 from docx import Document
 cfg=copy.deepcopy(CONFIG);cfg['custom_metrics']=[M];cfg['funds'].append(F)
 results=[]
 for f in cfg['funds']:
  for metric in metric_labels(cfg):results.append(resolve_metric([],f,metric,cfg))
 selected=next(x for x in results if x['ticker']=='TEST' and x['metric']==M['id'])
 c=extract_custom(BODY,F,FILING,SOURCE,[M])[0]
 selected.update(value=24.5,period='2024-12-31',sources=[c],reason='Reported portfolio turnover.',basis='custom:custom_turnover')
 run={'config':cfg,'results':results,'nav_history':[],'quarters':['2024-12-31'],'run_id':'test-run','judgment_mode':'rules_only','apex':{'status':'withheld','values':{}},'errors':[]}
 path=tmp_path/'benchmark.docx';export_word(run,path);doc=Document(path)
 text=' '.join(p.text for p in doc.paragraphs)+' '.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
 assert 'TEST' in text and 'Portfolio turnover' in text and '24.50%' in text
 assert 'Apex Ridge internal figures remain withheld' in text
 assert any(rel.target_ref==SOURCE['url'] for rel in doc.part.rels.values() if rel.reltype.endswith('/hyperlink'))
 assert len(doc.tables)==2  # concise attachment; history remains in detailed evidence
 assert 'Institutional' in text and 'Class I' in text
 assert "position relative to the competitor set cannot yet be assessed" in text
 assert 'Definitions and source notes' in text and '[1]' in text
 detailed=tmp_path/'benchmark-evidence.docx';export_evidence_word(run,detailed)
 detail=Document(detailed)
 assert any(p.text=='Evidence and review notes' for p in detail.paragraphs)
 assert len(detail.tables)>=3


def test_summary_does_not_invent_relative_performance_after_source_approval():
 run={'config':CONFIG,'results':[], 'apex':{'status':'approved','values':{'return_1y':{'value':99}}}}
 summary=committee_summary(run)
 assert 'relative positioning still requires PM confirmation' in summary
 assert 'outperform' not in summary and '99' not in summary


def test_workspace_run_snapshots_settings_and_reuses_sources(tmp_path,monkeypatch):
 import time
 (tmp_path/'config').mkdir();(tmp_path/'config/engagement.json').write_text(json.dumps(CONFIG))
 report=tmp_path/'prior';report.mkdir();(report/'audit.json').write_text(json.dumps({'config':CONFIG}))
 w=Workspace(tmp_path,tmp_path/'config/workspace.json',report)
 w.save(dict(CONFIG,custom_metrics=[M]));calls=[]
 def run(args,**kwargs):
  calls.append(args)
  snapshot=json.loads(Path(args[args.index('--config')+1]).read_text())
  assert snapshot['custom_metrics']==[M]
  output=Path(args[args.index('--output')+1]);output.mkdir();(output/'index.html').write_text('review');(output/'benchmark.docx').write_bytes(b'word')
  return type('Completed',(),{'returncode':2})()
 monkeypatch.setattr('apex_benchmark.workspace.subprocess.run',run)
 w.start('cached')
 deadline=time.monotonic()+2
 while w.state()['job']['status']=='running' and time.monotonic()<deadline:time.sleep(.01)
 assert '--offline' in calls[0] and '--from-audit' in calls[0]
 assert w.state()['job']['status']=='review_required' and w.state()['has_word']


def test_custom_native_pdf_preserves_page_locator():
 from io import BytesIO
 from reportlab.pdfgen import canvas
 out=BytesIO();pdf=canvas.Canvas(out);pdf.drawString(50,700,'Portfolio turnover rate: 24.5% for the year ended December 31, 2024.');pdf.save()
 body=out.getvalue();source=dict(SOURCE,url='https://www.sec.gov/report.pdf',sha256=hashlib.sha256(body).hexdigest())
 c=extract_custom(body,F,FILING,source,[M])[0]
 assert c['locator']=='PDF page 1; native text' and '24.5%' in c['quote']
 assert not c['eligible'] and c['value'] is None
